import json,tempfile,unittest
from pathlib import Path
from docx import Document
from bridge_agent.documents import parse,search,save_json
from bridge_agent.extraction import validate_findings,chunks,source_text
from bridge_agent.audit import label_data,FIELDS
from bridge_agent.render import render,materialize
from bridge_agent.render import normalize_date
from bridge_agent.agent import fact_candidates
from bridge_agent.extraction import extract_all,FINDINGS
from unittest.mock import patch
from bridge_agent.model import Model

class ContractTests(unittest.TestCase):
    def test_initial_source_blocks_are_read_evidence_and_unknown_ids_stay_rejected(self):
        from bridge_agent.agent import synthesize,validate_report
        basic={k:{'value':'无','sources':[]} for k in FIELDS[:12]};basic['桥梁名称']={'value':'测试桥','sources':['B000001']}
        claim={'text':'原文未提供详细资料','sources':['B000002'],'inferred':False}
        report={'basic':basic,'has_history':False,'trend':claim,'overall':claim,'risks':claim,'conclusion':[claim]*3,'recommendation_ids':[],'causes':[claim],'safety':[claim]}
        doc={'filename':'test.docx','blocks':[{'id':'B000001','section':'','text':'测试桥'},{'id':'B000002','section':'','text':'原文未提供详细资料'}]}
        findings={k:[] for k in ['facts','recommendations','causes','safety','history_notes']};findings.update(diseases=[{'id':'D1','source':'B000001','part':'桥面','kind':'裂缝','description':'裂缝'}],coverage={'total_blocks':2,'processed_blocks':2,'batches':1})
        class FakeModel:
            def call(self,*args,**kwargs):return {'role':'assistant','tool_calls':[{'id':'call1','function':{'name':'submit_report','arguments':json.dumps(report)}}]}
        with tempfile.TemporaryDirectory() as td:
            self.assertEqual(synthesize(FakeModel(),doc,findings,Path(__file__).resolve().parents[1],Path(td)),report)
        with self.assertRaisesRegex(ValueError,r'overall.*B000002.*read_sources'):
            validate_report(report,doc,findings,{'B000001'})
        report['risks']={**claim,'sources':['B999999']}
        with self.assertRaisesRegex(ValueError,r'risks.*B999999'):
            validate_report(report,doc,findings,{'B000001','B000002'})
    def test_numbered_short_recommendations_are_not_lost(self):
        from bridge_agent.recommendations import action_candidates
        doc={'blocks':[{'id':'B1','kind':'row','section':'','cells':[],'text':'12 处理建议⑴按照规范对桥面坑洼部位进行修补；⑷恢复脱空支座；⑸粘贴桥墩反光标志；'}]}
        self.assertEqual([c['description'] for c in action_candidates(doc)],['⑴按照规范对桥面坑洼部位进行修补；','⑷恢复脱空支座；','⑸粘贴桥墩反光标志；'])
    def test_date_search_is_bounded_and_rejected_submission_can_read_missing_evidence(self):
        from bridge_agent.agent import synthesize
        basic={k:{'value':'无','sources':[]} for k in FIELDS[:12]};basic['桥梁名称']={'value':'测试桥','sources':['B000001']}
        claim={'text':'未提供历史资料','sources':['B000030'],'inferred':False}
        report={'basic':basic,'has_history':False,'trend':claim,'overall':claim,'risks':claim,'conclusion':[claim]*3,'recommendation_ids':[],'causes':[claim],'safety':[claim]}
        doc={'filename':'test.docx','blocks':[{'id':f'B{i:06d}','section':'','text':'测试桥，未提供历史资料'} for i in range(1,31)]}
        findings={k:[] for k in ['facts','recommendations','causes','safety','history_notes']};findings.update(diseases=[{'id':'D1','source':'B000001','part':'桥面','kind':'裂缝','description':'裂缝'}],coverage={'total_blocks':30,'processed_blocks':30,'batches':1})
        choices=[]
        class FakeModel:
            def call(self,messages,tools,choice,**kwargs):
                choices.append(choice);turn=len(choices)
                if turn<=6:name='search_document';args={'query':f'报告日期 检索{turn}'}
                elif turn==8:name='read_sources';args={'ids':['B000030']}
                else:name='submit_report';args=report
                return {'role':'assistant','tool_calls':[{'id':f'c{turn}','function':{'name':name,'arguments':json.dumps(args)}}]}
        with tempfile.TemporaryDirectory() as td:
            self.assertEqual(synthesize(FakeModel(),doc,findings,Path(__file__).resolve().parents[1],Path(td)),report)
            trace=json.loads((Path(td)/'agent_trace.json').read_text())
        self.assertIn('检索两次',trace[2]['result']['error']);self.assertIn('B000030',trace[6]['result']['error'])
        self.assertEqual(choices[6]['function']['name'],'submit_report');self.assertEqual(choices[7],'auto');self.assertEqual(choices[8]['function']['name'],'submit_report')
    def test_grouped_summary_rejects_invented_totals_and_keeps_all_members(self):
        from bridge_agent.selection import materialize_group
        group=[{'id':f'D{i}','source':f'B{i}','part':'箱梁','kind':'裂缝','description':f'位置{i}裂缝W=0.1mm','table_context':'表1 左幅主桥箱梁内'} for i in range(1,5)]
        spec={'part':'左幅主桥箱梁','kind':'裂缝','summary_source':'B9','summary_quote':'存在4条裂缝','example_ids':['D1','D2']}
        got=materialize_group(spec,group,{'B9':{'text':'左幅主桥存在4条裂缝'}})
        self.assertEqual(got['member_count'],4);self.assertEqual(got['member_ids'],['D1','D2','D3','D4']);self.assertIn('B4',got['sources']);self.assertIn('存在4条裂缝',got['description']);self.assertIn('左幅',got['part']);self.assertIn('主桥',got['part'])
        spec['part']='右幅主桥箱梁'
        with self.assertRaises(ValueError):materialize_group(spec,group,{'B9':{'text':'左幅主桥存在4条裂缝'}})
        spec['part']='箱梁'
        spec['summary_quote']='存在5条裂缝'
        with self.assertRaises(ValueError):materialize_group(spec,group,{'B9':{'text':'左幅主桥存在4条裂缝'}})
    def test_small_structured_ledger_survives_model_omission(self):
        from bridge_agent.extraction import structured_candidates
        b={'id':'B000001','kind':'row','row':2,'table':1,'section':'检测','table_context':'左幅主桥','header':['病害位置','病害类型','病害描述'],'cells':['箱梁','裂缝','L=1m,W=0.1mm'],'text':'箱梁|裂缝|L=1m,W=0.1mm'}
        doc={'filename':'x.docx','tables':[[b['header'],b['cells']]],'blocks':[b]}
        def no_diseases(*args):return {k:[] for k in FINDINGS['properties']}
        with tempfile.TemporaryDirectory() as td,patch('bridge_agent.extraction.extract_batch',side_effect=no_diseases):
            got=extract_all(None,doc,Path(td));self.assertEqual(len(got['diseases']),1);self.assertEqual(got['diseases'][0]['table_context'],'左幅主桥')
    def test_compound_score_quote_restores_original_not_paraphrase(self):
        data={k:[] for k in FINDINGS['properties']};source='主桥总体评分83.2；引桥总体评分75.2。'
        data['facts']=[{'field':'总体评分','value':'主桥83.2；引桥75.2','source':'B1','quote':'主桥83.2，引桥75.2'}]
        got=validate_findings(data,[{'id':'B1','text':source,'section':''}]);self.assertEqual(got['facts'][0]['quote'],source);self.assertTrue(got['facts'][0]['quote_reconstructed'])
        data['facts'][0].update(value='主桥93.2；引桥75.2',quote='改写')
        with self.assertRaises(ValueError):validate_findings(data,[{'id':'B1','text':source,'section':''}])
    def test_legacy_merged_cell_recommendation_list_is_split_without_copying_cover(self):
        from bridge_agent.recommendations import action_candidates
        doc={'blocks':[{'id':'B1','kind':'row','section':'概况','text':'处理意见|⑴原摘要建议布置桥面排水设施；⑵原摘要建议修复桥面破损；','cells':['处理意见','⑴原摘要建议布置桥面排水设施；⑵原摘要建议修复桥面破损；']},{'id':'B2','kind':'row','section':'动力试验','text':'8 评估结论⑴技术状况良好。9 处理建议⑴按照规范布置桥面排水设施，提高耐久性；⑵修复桥面混凝土铺装破损部位；','cells':[]}]}
        got=action_candidates(doc);self.assertEqual(len(got),2);self.assertTrue(all(c['source']=='B2' for c in got));self.assertNotIn('评估结论',str(got));self.assertTrue(got[1]['description'].startswith('⑵'))
    def test_explicit_recommendation_actions_keep_source_text_and_all_ids(self):
        from bridge_agent.recommendations import action_candidates,preserve_actions
        blocks=[{'id':'B1','kind':'paragraph','section':'14.2 养护建议','text':'（1）严格控制车辆荷载，严禁超载车辆通行。'},{'id':'B2','kind':'paragraph','section':'14.2 养护建议','text':'（2）其它病害应参照规范的相关条款进行处理。'}]
        class FakeModel:
            def call(self,*args,**kwargs):return {'tool_calls':[{'function':{'arguments':json.dumps({'actions':[{'id':f'C{i:04d}','part':'桥梁','category':'预防性养护','reason':''} for i in [1,2]]})}}]}
        doc={'blocks':blocks};self.assertEqual(len(action_candidates(doc)),2);findings={'recommendations':[{'description':'其它章节建议'}]};preserve_actions(FakeModel(),doc,findings);self.assertEqual(len(findings['recommendations']),2);self.assertEqual(findings['recommendations'][1]['description'],blocks[1]['text']);self.assertEqual(len(findings['additional_recommendations']),1)
    def test_explicit_recommendation_omission_is_rejected(self):
        from bridge_agent.agent import validate_report
        basic={k:{'value':'无','sources':[]} for k in FIELDS[:12]};basic['桥梁名称']={'value':'测试桥','sources':['B1']};claim={'text':'未提供历史资料','sources':['B1'],'inferred':False}
        report={'basic':basic,'has_history':False,'trend':claim,'overall':claim,'risks':claim,'conclusion':[claim]*3,'recommendation_ids':['R1'],'causes':[claim],'safety':[claim]};doc={'blocks':[{'id':'B1','text':'测试桥'}]};findings={'facts':[],'diseases':[{'description':'裂缝'}],'recommendations':[{'id':'R1'},{'id':'R2'}],'recommendation_mode':'explicit_recommendation_chapter','coverage':{'processed_blocks':1,'total_blocks':1}}
        with self.assertRaises(ValueError):validate_report(report,doc,findings,{'B1'})
        report['recommendation_ids']=['R1','R2'];validate_report(report,doc,findings,{'B1'})
    def test_summary_segment_id_restores_original_without_model_transcription(self):
        from bridge_agent.selection import materialize_group
        rows=[{'id':'D1','source':'B1','part':'箱梁','kind':'裂缝','description':'位置1裂缝'},{'id':'D2','source':'B2','part':'箱梁','kind':'裂缝','description':'位置2裂缝'}]
        spec={'part':'箱梁','kind':'裂缝','summary_source':'B9:0','summary_quote':'','example_ids':['D1']}
        got=materialize_group(spec,rows,{'B9:0':{'source':'B9','text':'存在2条裂缝；'}});self.assertEqual(got['summary_source'],'B9');self.assertEqual(got['summary_quote'],'存在2条裂缝；');self.assertIn('B9',got['sources']);self.assertNotIn('B9:0',got['sources'])
    def test_group_organizer_repairs_plain_text_with_explicit_tool_instruction(self):
        from bridge_agent.selection import selection
        class FakeModel:
            concurrency=1
            def __init__(self):self.calls=0
            def call(self,messages,*args,**kwargs):
                self.calls+=1
                if self.calls==1:return {'role':'assistant','content':'已整理'}
                self.asserted='organize_groups' in messages[-1]['content']
                return {'role':'assistant','tool_calls':[{'id':'c1','function':{'arguments':json.dumps({'groups':[{'id':'G0000','part':'桥墩','kind':'裂缝','summary_source':'','summary_quote':'','example_ids':['D0']}]})}}]}
        doc={'filename':'未知测试桥.docx','blocks':[{'id':'B000001','kind':'row','table':1,'section':'检查','text':'裂缝'}]};findings={'facts':[],'diseases':[{'id':f'D{i}','source':'B000001','part':'桥墩','kind':'裂缝','description':f'位置{i}裂缝'} for i in range(81)]};model=FakeModel()
        with tempfile.TemporaryDirectory() as td:
            got=selection(model,doc,findings,Path(td)/'selection.json');self.assertEqual(model.calls,2);self.assertTrue(model.asserted);self.assertEqual(len(got['rows']),1);self.assertEqual(got['rows'][0]['description'],'位置0裂缝')
    def test_style_profile_uses_other_year_and_excludes_validation_labels(self):
        from bridge_agent.selection import writing_style
        a,p=writing_style({'filename':'2、马桑溪大桥.docx'},{'facts':[{'field':'报告日期','value':'2018年12月31日'}]});self.assertEqual(a,'representative');self.assertEqual(p['year'],2017)
        a,p=writing_style({'filename':'4、高家花园.docx'},{'facts':[{'field':'报告日期','value':'2018年12月31日'}]});self.assertEqual(a,'aggregate');self.assertEqual(p['year'],2017)
        root=Path(__file__).resolve().parents[1];profiles=json.loads((root/'assets/style_profiles.json').read_text());self.assertFalse(set(x['label'] for x in profiles['profiles'])&set(profiles['excluded_labels']))
    def test_main_bridge_overall_score_is_not_superstructure_score(self):
        from bridge_agent.extraction import validate_metric_scope
        with self.assertRaises(ValueError):validate_metric_scope('上部结构评分','该桥主桥总体技术状况评定为83.2，等级为二类')
        validate_metric_scope('总体评分','该桥主桥总体技术状况评定为83.2，等级为二类')
        validate_metric_scope('上部结构评分','该桥主桥上部结构评分为83.2')
    def test_unknown_history_can_omit_empty_quote_without_crashing(self):
        from bridge_agent.history import match_history
        class FakeModel:
            def call(self,*args,**kwargs):
                return {'tool_calls':[{'function':{'arguments':json.dumps({'matches':[{'id':'D1','new':'无法判断','previous':'无','development':'无','sources':[]}]})}}]}
        doc={'blocks':[{'id':'B1','section':'历年对比','kind':'paragraph','text':'历史未提供详细位置'}]};findings={'history_notes':[],'diseases':[{'id':'D1','source':'B2','part':'1#墩','kind':'裂缝','description':'裂缝'}]}
        with tempfile.TemporaryDirectory() as td:
            got=match_history(FakeModel(),doc,findings,Path(td));self.assertEqual(got['matched'],0);self.assertEqual(findings['diseases'][0]['history_quote'],'')
    def test_rating_table_does_not_duplicate_detailed_diseases(self):
        bs=[{'id':'B000001','kind':'row','row':2,'table':1,'section':'检测','header':['病害位置','病害类型','病害描述'],'cells':['桥墩','涂层剥落','墩身涂层剥落'],'text':'墩身涂层剥落'}, {'id':'B000002','kind':'row','row':2,'table':2,'section':'养护建议','header':['部位','损坏类型','扣分值'],'cells':['墩身','涂层剥落','12'],'text':'墩身|涂层剥落|12'}]
        doc={'filename':'x.docx','tables':[[bs[0]['header'],bs[0]['cells']],[bs[1]['header'],bs[1]['cells']]],'blocks':bs}
        def fake(*args):
            data={k:[] for k in FINDINGS['properties']};data['diseases']=[{'source':'B000002','part':'墩身','kind':'涂层剥落','description':'涂层剥落','history_sources':[]}];return data
        with tempfile.TemporaryDirectory() as td,patch('bridge_agent.extraction.extract_batch',side_effect=fake):
            got=extract_all(None,doc,Path(td));self.assertEqual(len(got['diseases']),1);self.assertEqual(got['diseases'][0]['source'],'B000001')
    def test_caption_scope_persists_across_chunk_boundaries(self):
        from bridge_agent.extraction import pack_blocks
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);p=root/'scope.docx';d=Document();d.add_paragraph('表6.1-2 左幅主桥箱梁外检查结果');t=d.add_table(rows=2,cols=2);t.cell(0,0).text='位置';t.cell(0,1).text='病害';t.cell(1,0).text='第9跨';t.cell(1,1).text='露筋';d.save(p)
            doc=parse(p,root/'cache');last=doc['blocks'][-1];self.assertIn('左幅主桥',last['table_context']);self.assertIn('左幅主桥',json.dumps(pack_blocks([last]),ensure_ascii=False))
    def test_main_bridge_overview_preserves_approach_facts_and_precision(self):
        basic={k:{'value':'无','sources':[]} for k in FIELDS[:12]};basic['总体评分']['value']='主桥70.80；引桥74.00';basic['上一次总体评分']['value']='主桥74；引桥68.6';basic['总体等级']['value']='主桥二类；引桥二类'
        report={'basic':basic,'has_history':True,'recommendation_ids':[],'trend':{'text':'趋势'},'overall':{'text':'结论'},'risks':{'text':'风险'},'conclusion':[],'causes':[],'safety':[]};findings={'recommendations':[],'diseases':[]}
        result=materialize(report,findings);self.assertEqual(result['summary']['总体评分'],'70.80');self.assertEqual(result['summary']['上一次总体评分'],'74.00');self.assertEqual(result['summary']['总体等级'],'二类');self.assertIn('引桥',basic['总体评分']['value'])
        basic['上一次总体评分']['value']='主桥74.234；引桥68.6';self.assertEqual(materialize(report,findings)['summary']['上一次总体评分'],'74.234')
    def test_upstream_failure_splits_only_the_failed_source_batch(self):
        def fake_extract(model,blocks,path,identity):
            if len(blocks)>1:raise RuntimeError('Model HTTP 502: upstream failure')
            return {k:[] for k in FINDINGS['properties']}
        doc={'filename':'x.docx','tables':[],'blocks':[{'id':'B000001','text':'第一段','section':''},{'id':'B000002','text':'第二段','section':''}]}
        with tempfile.TemporaryDirectory() as td,patch('bridge_agent.extraction.extract_batch',side_effect=fake_extract) as mocked:
            result=extract_all(None,doc,Path(td));self.assertEqual(result['coverage']['processed_blocks'],2);self.assertEqual(mocked.call_count,3);self.assertEqual(json.loads((Path(td)/'0000.json').read_text())['split_depth'],1)
    def test_shared_response_requires_identical_request_and_service(self):
        class Response:
            status_code=200
            def json(self):return {'choices':[{'finish_reason':'tool_calls','message':{'role':'assistant','tool_calls':[{'id':'call1','type':'function','function':{'name':'done','arguments':'{}'}}]}}],'usage':{'prompt_tokens':5,'completion_tokens':2}}
        class Client:
            def __init__(self,forbidden=False):self.forbidden=forbidden
            def post(self,*a,**k):
                if self.forbidden:raise AssertionError('Unexpected network call')
                return Response()
        with tempfile.TemporaryDirectory() as td,patch.dict('os.environ',{'INSPECTION_FRESH_RUN':'0'}):
            root=Path(td);env=root/'.env';env.write_text('INSPECTION_API_BASE_URL=https://model.invalid/v1\nINSPECTION_API_KEY=dummy\nINSPECTION_MODEL=test-model\n')
            a=Model(root,root/'work/first');a.client.close();a.client=Client();messages=[{'role':'user','content':'source A'}];result=a.call(messages,[])
            b=Model(root,root/'work/second');b.client.close();b.client=Client(True);self.assertEqual(b.call(messages,[]),result);self.assertTrue(b.records[-1]['cache'])
            env.write_text(env.read_text().replace('model.invalid','other.invalid'));c=Model(root,root/'work/third');c.client.close();c.client=Client();c.call(messages,[]);self.assertFalse(c.records[-1]['cache'])
    def test_nested_table_is_not_flattened_into_parent_row(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);p=root/'nested.docx';d=Document();outer=d.add_table(rows=1,cols=1);cell=outer.cell(0,0);cell.text='前文';inner=cell.add_table(rows=2,cols=2);inner.cell(0,0).text='部位';inner.cell(0,1).text='病害描述';inner.cell(1,0).text='3#墩';inner.cell(1,1).text='裂缝W=0.12mm';cell.add_paragraph('后文');d.save(p)
            parsed=parse(p,root/'cache');texts=[b['text'] for b in parsed['blocks']]
            self.assertEqual(parsed['table_count'],2);self.assertEqual(sum('0.12mm' in t for t in texts),1);self.assertLess(texts.index('前文'),next(i for i,t in enumerate(texts) if '0.12mm' in t));self.assertEqual(texts[-1],'后文')
    def test_date_and_consensus_normalization(self):
        self.assertEqual(normalize_date('二〇一八年十二月三十一日'),'2018年12月31日');self.assertEqual(normalize_date('2014/7/25'),'2014年7月25日')
        rows=[{'field':'总体评分','value':v,'source':s,'quote':'原文'} for v,s in [('70.80','B1'),('70.8','B2'),('74','B3'),('70.8','B1')]]
        got=fact_candidates(rows);self.assertEqual(len(got),2);self.assertEqual(got[0]['sources'],['B1','B2'])
    def test_parser_order_table_and_unicode(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);p=root/'fake.doc';d=Document();d.add_paragraph('1、桥梁概况');t=d.add_table(rows=2,cols=2);t.cell(0,0).text='位置';t.cell(0,1).text='描述';t.cell(1,0).text='3#墩';t.cell(1,1).text='裂缝W=0.12mm';d.add_paragraph('结束');d.save(p)
            parsed=parse(p,root/'cache');self.assertEqual(parsed['format'],'docx');self.assertEqual([b['kind'] for b in parsed['blocks']],['paragraph','row','row','paragraph']);self.assertIn('0.12mm',search(parsed['blocks'],'裂缝')[0]['text']);self.assertEqual(parsed['blocks'][2]['cells'],['3#墩','裂缝W=0.12mm'])
            again=parse(p,root/'cache');self.assertEqual(parsed['sha256'],again['sha256'])
            d.add_paragraph('新增信息');d.save(p);self.assertNotEqual(parsed['sha256'],parse(p,root/'cache')['sha256'])
    def test_source_columns_and_fabricated_quote(self):
        b={'id':'B000001','kind':'row','cells':['3#墩','裂缝W=0.12mm'],'text':'3#墩 | 裂缝W=0.12mm','section':''}
        self.assertEqual(source_text({'source':b['id'],'cols':[1]}, {b['id']:b}),'裂缝W=0.12mm')
        with self.assertRaises(ValueError):source_text({'source':b['id'],'cols':[],'quote':'裂缝W=1.2mm'}, {b['id']:b})
        with self.assertRaises(ValueError):source_text({'source':b['id'],'cols':[5]}, {b['id']:b})
    def test_numeric_fact_and_quote_restore(self):
        data={k:[] for k in ['facts','diseases','recommendations','causes','safety','history_notes','no_disease_sources','excluded_diseases']};b={'id':'B1','text':'整体BCI=89.96，为B级','section':''}
        data['facts']=[{'field':'总体评分','value':'89.96','source':'B1','quote':'改写的句子'}]
        got=validate_findings(data,[b]);self.assertTrue(got['facts'][0]['quote_reconstructed']);self.assertIn(got['facts'][0]['quote'],b['text'])
        data['facts'][0].update(value='99.99',quote=b['text'])
        with self.assertRaises(ValueError):validate_findings(data,[b])
    def test_character_budget_full_coverage(self):
        bs=[{'id':f'B{i}','text':'裂缝'*100,'section':'章节'} for i in range(30)];gs=chunks(bs,1000)
        self.assertEqual([b['id'] for g in gs for b in g],[b['id'] for b in bs]);self.assertGreater(len(gs),1)
        with self.assertRaises(ValueError):chunks([{'id':'B1','text':'x'*10000,'section':''}],500)
    def test_html_label_without_header(self):
        d={'tables':[[['桥梁名称','某桥','全称'],['总体评分','90','评分']],[['序号','位置','病害种类','病害情况','是否新增','上一次定检状态','发展程度'],['1','3#墩','裂缝','W=0.1mm','否','无','无']]]}
        result=label_data(d);self.assertEqual(result['summary']['桥梁名称'],'某桥');self.assertEqual(len(result['diseases']),1)
    def test_docx_full_contract_readback(self):
        root=Path(__file__).resolve().parent.parent
        basic={k:{'value':'无','sources':[]} for k in FIELDS[:12]};basic['桥梁名称']['value']='测试桥'
        claim=lambda s:{'text':s,'sources':['B1'],'inferred':False}
        report={'basic':basic,'has_history':False,'trend':claim('未提供历史资料'),'overall':claim('裂缝需处置'),'risks':claim('3#墩裂缝'),'conclusion':[claim('经综合评定，无评分。'),claim('无历史资料。'),claim('目前存在裂缝。')],'recommendation_ids':['R1'],'causes':[claim('原文未说明')],'safety':[claim('可能影响耐久性')]}
        findings={'recommendations':[{'id':'R1','category':'尽快维修','description':'（1）封闭裂缝','part':'3#墩'}],'diseases':[{'part':'3#墩','kind':'裂缝','description':'裂缝W=0.12mm','source':'B1'}]}
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'out.docx';data=render(p,report,findings,root);doc=Document(p)
            self.assertEqual(doc.tables[1].rows[1].cells[2].text,'封闭裂缝');self.assertEqual([len(t.columns) for t in doc.tables],[3,4,7]);self.assertEqual(len(doc.tables[0].rows),17);self.assertEqual(data['diseases'][0][-3:],['否','无','无']);self.assertIn('1条尽快维修',data['summary']['建议'])
if __name__=='__main__':unittest.main()

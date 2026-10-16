"""Check final DOCX structure, candidate membership and exact source spans; no API calls."""
import json,re,sys
from collections import Counter
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from bridge_agent.documents import parse,save_json
from bridge_agent.pipeline import code_hash,contract_hash
from scripts.validate import SAMPLES


def main():
 root=Path(__file__).resolve().parents[1]
 matrix=json.loads((root/'artifacts/final_delivery.json').read_text())
 cold=json.loads((root/'artifacts/final_delivery_cold.json').read_text())
 overrides={x['sample']:x for x in cold};results=[]
 sources={name:root/source for name,source,_ in SAMPLES}
 for original in matrix:
  item=overrides.get(original['sample'],original);m=item['metrics']
  assert item['status']=='succeeded'
  assert m['code_hash']==code_hash() and m['contract_hash']==contract_hash(root)
  run='final_delivery_cold' if item['sample'] in overrides else 'final_delivery'
  task=root/'work'/run/m['job_id']
  findings=json.loads((task/'findings.json').read_text())
  choice=json.loads((task/'selection.json').read_text())
  analysis=json.loads((task/'analysis.json').read_text())
  doc=parse(sources[item['sample']],root/'work/cache')
  blocks={b['id']:b['text'] for b in doc['blocks']}
  groups=choice.get('organized_groups',findings['diseases'])
  members=[i for g in groups for i in g.get('member_ids',[g['id']])]
  assert Counter(members)==Counter(d['id'] for d in findings['diseases'])
  for group in groups:
   assert all(s in blocks for s in group.get('sources',[group['source']]))
   if group.get('summary_quote'):assert group['summary_quote'] in blocks[group['summary_source']]
  required={r['id'] for r in findings['recommendations'] if r.get('canonical')}
  assert required.issubset(analysis['recommendation_ids'])
  for recommendation in findings['recommendations']:
   assert recommendation['description'] in blocks[recommendation['source']]
  word=Document(m['output']);columns=[len(t.columns) for t in word.tables]
  assert columns==[3,4,7] and len(word.tables[0].rows)==17
  assert len(word.tables[1].rows)-1==m['recommendations']
  assert len(word.tables[2].rows)-1==m['diseases']
  text='\n'.join([p.text for p in word.paragraphs]+[c.text for t in word.tables for row in t.rows for c in row.cells])
  assert not re.search(r'\bB\d{6}\b',text)
  results.append({'sample':item['sample'],'run':run,'code_hash':m['code_hash'],'contract_hash':m['contract_hash'],'docx_columns':columns,'summary_fields':16,'published_diseases':m['diseases'],'complete_group_members':len(members),'canonical_recommendations':len(required),'all_group_source_ids_valid':True,'exact_summary_source_spans':True,'recommendation_source_spans_and_submission_complete':True,'internal_source_ids_absent_from_word':True,'scope':'structural and ledger checks only; not semantic approval'})
 save_json(root/'artifacts/structural_and_ledger_checks.json',results)
 print(json.dumps(results,ensure_ascii=False,indent=2))


if __name__=='__main__':main()

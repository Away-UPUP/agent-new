from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .audit import FIELDS,REC_HEADERS,DISEASE_HEADERS
from .documents import read_json
import re
from decimal import Decimal

def normalize_date(value):
    digits=dict(zip('零〇○一二三四五六七八九','000123456789'))
    def number(s):
        if s.isdigit():return str(int(s))
        if '十' in s:
            a,b=s.split('十',1);return str((int(digits.get(a,'1')) if a else 1)*10+(int(digits[b]) if b else 0))
        return ''.join(digits.get(c,c) for c in s)
    pattern=r'([\d零〇○一二三四五六七八九]{4})[年./-]([\d一二三四五六七八九十]{1,3})(?:[月./-]([\d一二三四五六七八九十]{1,3})日?)?月?'
    def replace(m):return number(m[1])+'年'+number(m[2])+'月'+(number(m[3])+'日' if m[3] else '')
    return re.sub(pattern,replace,value)


def materialize(report,findings):
    index={r['id']:r for r in findings['recommendations']};recs=[index[i] for i in report['recommendation_ids']]
    counts=Counter(r['category'] for r in recs)
    summary={k:report['basic'][k]['value'] for k in FIELDS[:12]}
    # Labels use the main bridge in overview fields; detailed conclusions retain approach-bridge facts.
    scope=re.match(r'^主桥[：:\s]*(\d+(?:\.\d+)?)',summary['总体评分'])
    precision=len(scope[1].split('.')[1]) if scope and '.' in scope[1] else None
    for field in ['总体评分','总体等级','上一次总体评分','上一次总体等级']:
        match=re.match(r'^主桥[：:\s]*([^；;\n]+)',summary[field])
        if match:
            summary[field]=match[1].strip()
            if '评分' in field and precision is not None and re.fullmatch(r'\d+(?:\.\d+)?',summary[field]):
                places=max(precision,len(summary[field].partition('.')[2]));summary[field]=f'{Decimal(summary[field]):.{places}f}'
    summary['报告日期']=normalize_date(summary['报告日期'])
    for field in [k for k in FIELDS[:12] if '等级' in k]:
        if re.fullmatch('[A-Ea-e]',summary[field]):summary[field]=summary[field].upper()+'级'
    summary.update({'病害发展趋势与具体说明':report['trend']['text'],'总体结论':report['overall']['text'],'主要风险点':report['risks']['text'],'建议':'、'.join(f'{counts[c]}条{c}' for c in ['立即处置','尽快维修','预防性养护'])})
    diseases=[]
    for i,d in enumerate(findings['diseases'],1):
        no_history=not report['has_history'];new='否' if no_history else d.get('new') or '无法判断';prev='无' if no_history else d.get('previous') or '无';development='无' if no_history else d.get('development') or '无'
        diseases.append([str(i),d['part'],d['kind'],d['description'],new,prev,development])
    return {'summary':summary,'recommendations':[[str(i),r['category'],re.sub(r'^\s*(?:[（(][一二三四五六七八九十\d]+[）)]|[⑴-⒇①-⑳])\s*','',r['description']),r['part']] for i,r in enumerate(recs,1)],'diseases':diseases,'conclusion':[c['text'] for c in report['conclusion']],'causes':[c['text'] for c in report['causes']],'safety':[c['text'] for c in report['safety']]}


def render(output,report,findings,root):
    data=materialize(report,findings);spec=read_json(Path(root)/'assets/label_spec.json')
    doc=Document(Path(root)/'assets/layout.docx')
    style=doc.styles['Normal'];style.font.name='宋体';style.font.size=Pt(11);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    style.paragraph_format.space_after=Pt(0);style.paragraph_format.line_spacing=1.0
    # Modal margins measured in 34/44 sampled labels, with source template page geometry.
    for sec in doc.sections:sec.top_margin=Inches(1);sec.bottom_margin=Inches(1);sec.left_margin=Inches(1.25);sec.right_margin=Inches(1.25)
    def para(text,bold=False):
        p=doc.add_paragraph();r=p.add_run(text);r.bold=bold;r.font.name='宋体';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');return p
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers))
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'000000');borders.append(el)
        t._tbl.tblPr.append(borders)
        for c,text in zip(t.rows[0].cells,headers):c.text=text
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0)
                    for r in p.runs:r.font.size=Pt(10.5);r.font.name='宋体';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        if widths:
            t.autofit=False
            sec=doc.sections[0];available=(sec.page_width-sec.left_margin-sec.right_margin)/360000
            widths=[w*available/sum(widths) for w in widths]
            for col,width in zip(t.columns,widths):col.width=Cm(width)
            for row in t.rows:
                for c,width in zip(row.cells,widths):c.width=Cm(width)
        return t
    para('1、简要信息',True)
    table(spec['summary_headers'],[[k,data['summary'][k],spec['descriptions'].get(k,'')] for k in FIELDS],[3.3,7.5,4.2])
    para('2、详细信息',True);para('（1）详细结论',True)
    for s in data['conclusion']:para(s)
    para('（2）建议明细',True);table(REC_HEADERS,data['recommendations'],[0.9,2.0,9.6,2.5])
    para('（3）病害列表',True);table(DISEASE_HEADERS,data['diseases'],[0.8,2.1,1.5,5.5,1.1,2.8,1.2])
    para('病害成因：',True)
    for i,s in enumerate(data['causes'],1):para(f'（{i}）{s}')
    para('处置建议：',True)
    for i,r in enumerate(data['recommendations'],1):para(f'（{i}）{r[2]}')
    para('安全影响：',True)
    for s in data['safety']:para(s)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);tmp=output.with_name(output.stem+'.tmp.docx');doc.save(tmp)
    check=Document(tmp)
    if len(check.tables)!=3 or len(check.tables[0].rows)!=17 or len(check.tables[2].rows)!=len(data['diseases'])+1:raise ValueError('DOCX readback structure mismatch')
    if [c.text for c in check.tables[2].rows[0].cells]!=DISEASE_HEADERS:raise ValueError('Disease table header mismatch')
    for row,expected in zip(check.tables[2].rows[1:],data['diseases']):
        if [c.text for c in row.cells]!=expected:raise ValueError('Disease text changed during Word export')
    tmp.replace(output);return data
